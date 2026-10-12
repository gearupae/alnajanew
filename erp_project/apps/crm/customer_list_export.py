"""PDF and Excel exports for the CRM customer list."""
from __future__ import annotations

import html
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.settings_app.models import CompanySettings

CUSTOMER_EXPORT_COLUMNS = [
    {'key': 'customer_number', 'label': '#', 'width': 22},
    {'key': 'company', 'label': 'Customer', 'width': 36},
    {'key': 'name', 'label': 'Contact person', 'width': 24},
    {'key': 'email', 'label': 'Email', 'width': 32},
    {'key': 'phone', 'label': 'Phone', 'width': 18},
    {'key': 'type', 'label': 'Type', 'width': 14},
    {'key': 'salesman', 'label': 'Salesman', 'width': 22},
    {'key': 'status', 'label': 'Status', 'width': 14},
    {'key': 'trn', 'label': 'VAT (TRN)', 'width': 18},
]


def _company_context():
    cs = CompanySettings.get_settings()
    return {
        'company_name': cs.company_name or 'Company',
        'address': cs.address or '',
    }


def customer_export_row(customer) -> dict:
    return {
        'customer_number': customer.customer_number or '',
        'company': customer.company or '',
        'name': customer.name or '',
        'email': customer.email or '',
        'phone': customer.phone or '',
        'type': customer.get_customer_type_display(),
        'salesman': customer.assigned_salesman_label,
        'status': customer.get_status_display(),
        'trn': customer.trn or '',
    }


def export_customer_list_xlsx(rows: list[dict], *, title: str, filters_text: str, generated_by: str) -> bytes:
    ctx = _company_context()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Customers'

    brand = '1e3a5f'
    header_fill = PatternFill('solid', fgColor=brand)
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    col_count = len(CUSTOMER_EXPORT_COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    c1 = ws.cell(row=1, column=1, value=ctx['company_name'])
    c1.font = Font(size=16, bold=True, color=brand)
    c1.alignment = Alignment(horizontal='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    c2 = ws.cell(row=2, column=1, value=title)
    c2.alignment = Alignment(horizontal='center', wrap_text=True)
    c2.font = Font(size=12, bold=True)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)
    meta = f'Generated {date.today():%d/%m/%Y} | By {generated_by} | {len(rows)} record(s)'
    if filters_text:
        meta += f' | Filters: {filters_text}'
    ws.cell(row=3, column=1, value=meta).alignment = Alignment(horizontal='center', wrap_text=True)

    start_row = 5
    for col_idx, col in enumerate(CUSTOMER_EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col['label'])
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row_idx, row in enumerate(rows, start=start_row + 1):
        for col_idx, col in enumerate(CUSTOMER_EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(col['key'], ''))
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col_idx, col in enumerate(CUSTOMER_EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = min(40, max(10, col['width'] // 2))

    ws.auto_filter.ref = f'A{start_row}:{get_column_letter(col_count)}{start_row + len(rows)}'
    ws.freeze_panes = f'A{start_row + 1}'

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_customer_list_pdf(rows: list[dict], *, title: str, filters_text: str, generated_by: str) -> bytes:
    ctx = _company_context()
    buf = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        buf,
        pagesize=page_size,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustTitle',
        parent=styles['Heading2'],
        alignment=1,
        fontSize=13,
        spaceAfter=4,
        textColor=colors.HexColor('#1e3a5f'),
    )
    meta_style = ParagraphStyle(
        'CustMeta',
        parent=styles['Normal'],
        fontSize=8,
        alignment=1,
        textColor=colors.HexColor('#4b5563'),
        spaceAfter=8,
    )
    cell_style = ParagraphStyle(
        'CustCell',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
    )
    header_blue = colors.HexColor('#1e3a5f')
    alt1 = colors.HexColor('#ffffff')
    alt2 = colors.HexColor('#f7f9fc')

    story = []
    story.append(Paragraph(html.escape(ctx['company_name']), title_style))
    story.append(Paragraph(html.escape(title), title_style))
    meta = f'Generated {date.today():%d/%m/%Y} · {html.escape(generated_by)} · {len(rows)} record(s)'
    if filters_text:
        meta += f'<br/>Filters: {html.escape(filters_text)}'
    story.append(Paragraph(meta, meta_style))
    story.append(Spacer(1, 4))

    headers = [col['label'] for col in CUSTOMER_EXPORT_COLUMNS]
    table_data = [headers]
    for row in rows:
        table_data.append([
            Paragraph(html.escape(str(row.get(col['key'], '') or '')), cell_style)
            for col in CUSTOMER_EXPORT_COLUMNS
        ])

    col_widths = [doc.width * (col['width'] / sum(c['width'] for c in CUSTOMER_EXPORT_COLUMNS))
                  for col in CUSTOMER_EXPORT_COLUMNS]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), header_blue),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [alt1, alt2]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(table)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(colors.HexColor('#6b7280'))
        canvas.drawString(10 * mm, 8 * mm, html.unescape(ctx['company_name'])[:80])
        canvas.drawRightString(page_size[0] - 10 * mm, 8 * mm, f'Page {canvas.getPageNumber()}')
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    pdf = buf.getvalue()
    buf.close()
    return pdf
