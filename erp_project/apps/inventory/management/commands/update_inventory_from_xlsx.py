"""
Update inventory item unit, price, and stock quantity from Excel workbook.

Reads Sheet17-style layout (or auto-detects a sheet with ITEM NAME, UNIT,
UNIT PRICE AED, QTY). Matches existing items by normalized name.

Usage:
    python manage.py update_inventory_from_xlsx --file path/to/file.xlsx
    python manage.py update_inventory_from_xlsx --file path/to/file.xlsx --sheet Sheet17
    python manage.py update_inventory_from_xlsx --file path/to/file.xlsx --dry-run
"""
from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from difflib import get_close_matches

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.inventory.models import Item, Stock, Warehouse

SHEET17_HEADERS = {
    'name': ('ITEM NAME', 'Item Name', 'ITEM DESCRIPTION', 'Item Description', 'ITEM', 'DESCRIPTION'),
    'unit': ('UNIT', 'Unit'),
    'price': ('UNIT PRICE AED', 'Unit Price (AED)', 'PRICE', 'Sales Price', 'U/PRICE', 'RATE'),
    'qty': ('QTY', 'Qty', 'CLOSING', 'CLOSING STOCK', 'Available', 'AVAILABLE'),
}


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        import math
        if isinstance(value, float) and math.isnan(value):
            return True
    except TypeError:
        pass
    return False


def _cell_str(value) -> str:
    if _is_blank(value):
        return ''
    return str(value).strip()


def _parse_decimal(value) -> Decimal | None:
    if _is_blank(value):
        return None
    try:
        return Decimal(str(value)).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        return None


def _normalize_unit(raw: str) -> str:
    unit = _cell_str(raw).lower().replace("'", '')
    mapping = {
        'no': 'pcs',
        'nos': 'pcs',
        'nos.': 'pcs',
        'ea': 'pcs',
        'each': 'pcs',
        'meter': 'm',
        'metre': 'm',
        'mtr': 'm',
        'length': 'length',
        'lot': 'lot',
        'box': 'box',
        'pkt': 'pkt',
        'doz': 'doz',
        'kg': 'kg',
        'units': 'units',
        'pcs': 'pcs',
        'roll': 'roll',
        'set': 'set',
    }
    return mapping.get(unit, unit or 'pcs')[:20]


def _normalize_name(value) -> str:
    text = _cell_str(value).lower()
    text = text.replace("''", '"')
    text = re.sub(r'\([^)]*\)', '', text)
    text = re.sub(r'[\-–—]', ' ', text)
    text = re.sub(r'[^\w\s/]', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def _header_map(headers: list[str]) -> dict[str, int] | None:
    normalized = {_cell_str(h): idx for idx, h in enumerate(headers) if _cell_str(h)}
    lower_map = {k.lower(): v for k, v in normalized.items()}

    def find(keys):
        for key in keys:
            if key in normalized:
                return normalized[key]
            lk = key.lower()
            if lk in lower_map:
                return lower_map[lk]
        return None

    name_idx = find(SHEET17_HEADERS['name'])
    unit_idx = find(SHEET17_HEADERS['unit'])
    price_idx = find(SHEET17_HEADERS['price'])
    qty_idx = find(SHEET17_HEADERS['qty'])
    if name_idx is None:
        return None
    return {
        'name': name_idx,
        'unit': unit_idx,
        'price': price_idx,
        'qty': qty_idx,
    }


def _detect_sheet(wb, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise CommandError(f'Sheet not found: {sheet_name}')
        return wb[sheet_name]

    best = None
    best_score = 0
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(max_row=6, values_only=True))
        for row in rows:
            headers = [_cell_str(v) for v in row]
            mapping = _header_map(headers)
            if not mapping:
                continue
            score = 1 + sum(1 for k in ('unit', 'price', 'qty') if mapping.get(k) is not None)
            if score > best_score:
                best_score = score
                best = (name, row, mapping)
    if not best:
        raise CommandError('No sheet with ITEM NAME / UNIT / price / qty columns found.')
    return wb[best[0]], best[1], best[2]


def load_update_rows(path: str, sheet_name: str | None = None) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name or True:
        detected = _detect_sheet(wb, sheet_name)
        if sheet_name:
            ws = detected
            rows = list(ws.iter_rows(values_only=True))
            header_row = next(
                (row for row in rows[:10] if _header_map([_cell_str(v) for v in row])),
                None,
            )
            if not header_row:
                wb.close()
                raise CommandError(f'Could not find header row in {sheet_name}')
            col_map = _header_map([_cell_str(v) for v in header_row])
        else:
            ws, header_row, col_map = detected
            rows = list(ws.iter_rows(values_only=True))
    else:
        wb.close()
        raise CommandError('No sheet detected')

    header_idx = rows.index(header_row)
    data = []
    for row in rows[header_idx + 1:]:
        if not row or not any(not _is_blank(v) for v in row):
            continue
        name = _cell_str(row[col_map['name']]) if col_map['name'] < len(row) else ''
        if not name:
            continue
        unit_val = row[col_map['unit']] if col_map.get('unit') is not None and col_map['unit'] < len(row) else ''
        price_val = row[col_map['price']] if col_map.get('price') is not None and col_map['price'] < len(row) else None
        qty_val = row[col_map['qty']] if col_map.get('qty') is not None and col_map['qty'] < len(row) else None
        data.append({
            'name': name,
            'unit': _normalize_unit(unit_val) if not _is_blank(unit_val) else None,
            'price': _parse_decimal(price_val),
            'qty': _parse_decimal(qty_val),
        })
    wb.close()
    return data


class Command(BaseCommand):
    help = 'Update inventory unit, selling price, and stock quantity from Excel workbook.'

    def add_arguments(self, parser):
        parser.add_argument('--file', type=str, required=True, help='Path to Excel workbook')
        parser.add_argument('--sheet', type=str, default='', help='Sheet name (default: auto-detect Sheet17-style)')
        parser.add_argument('--dry-run', action='store_true', help='Report matches without saving')
        parser.add_argument(
            '--warehouse',
            type=str,
            default='',
            help='Warehouse for stock qty updates (default: first active warehouse)',
        )

    def _resolve_warehouse(self, name: str) -> Warehouse:
        if name:
            warehouse = Warehouse.objects.filter(name__iexact=name, is_active=True).first()
            if not warehouse:
                raise CommandError(f'Warehouse not found: {name}')
            return warehouse
        warehouse = Warehouse.objects.filter(is_active=True).order_by('pk').first()
        if not warehouse:
            raise CommandError('No active warehouse found.')
        return warehouse

    def _build_name_index(self):
        index: dict[str, list[Item]] = {}
        keys: list[str] = []
        for item in Item.objects.filter(is_active=True):
            key = _normalize_name(item.name)
            index.setdefault(key, []).append(item)
            keys.append(key)
        return index, keys

    def _find_item(self, row_name: str, index: dict[str, list[Item]], keys: list[str]):
        key = _normalize_name(row_name)
        matches = index.get(key, [])
        if len(matches) == 1:
            return matches[0], 'exact'
        if len(matches) > 1:
            return None, 'ambiguous'

        fuzzy = get_close_matches(key, keys, n=2, cutoff=0.9)
        if len(fuzzy) == 1:
            return index[fuzzy[0]][0], 'fuzzy'
        return None, 'no_match'

    def handle(self, *args, **options):
        path = options['file']
        sheet = (options.get('sheet') or '').strip() or None
        dry_run = options['dry_run']

        try:
            rows = load_update_rows(path, sheet)
        except FileNotFoundError as exc:
            raise CommandError(str(exc)) from exc

        if not rows:
            raise CommandError('No rows found in workbook.')

        warehouse = None if dry_run else self._resolve_warehouse(options['warehouse'])
        name_index, name_keys = self._build_name_index()

        stats = {
            'rows': len(rows),
            'matched': 0,
            'matched_fuzzy': 0,
            'updated': 0,
            'stock_updated': 0,
            'skipped_no_match': 0,
            'skipped_ambiguous': 0,
            'skipped_no_changes': 0,
        }
        unmatched_samples = []

        @transaction.atomic
        def run():
            for row in rows:
                item, match_kind = self._find_item(row['name'], name_index, name_keys)
                if match_kind == 'no_match':
                    stats['skipped_no_match'] += 1
                    if len(unmatched_samples) < 15:
                        unmatched_samples.append(row['name'])
                    continue
                if match_kind == 'ambiguous' or item is None:
                    stats['skipped_ambiguous'] += 1
                    continue

                stats['matched'] += 1
                if match_kind == 'fuzzy':
                    stats['matched_fuzzy'] += 1
                changed_fields = []
                update_fields = []

                if row['unit'] and row['unit'] != item.unit:
                    item.unit = row['unit']
                    changed_fields.append('unit')
                    update_fields.append('unit')

                if row['price'] is not None and row['price'] != item.selling_price:
                    item.selling_price = row['price']
                    changed_fields.append('price')
                    update_fields.append('selling_price')
                    if item.purchase_price in (None, Decimal('0.00')):
                        item.purchase_price = row['price']
                        update_fields.append('purchase_price')

                if changed_fields:
                    if not dry_run:
                        item.save(update_fields=update_fields)
                    stats['updated'] += 1
                else:
                    stats['skipped_no_changes'] += 1

                if row['qty'] is not None and warehouse:
                    stock, _created = Stock.objects.get_or_create(
                        item=item,
                        warehouse=warehouse,
                        defaults={'quantity': row['qty']},
                    )
                    if stock.quantity != row['qty']:
                        stock.quantity = row['qty']
                        if not dry_run:
                            stock.save(update_fields=['quantity'])
                        stats['stock_updated'] += 1

            if dry_run:
                transaction.set_rollback(True)

        run()

        prefix = '[DRY RUN] ' if dry_run else ''
        self.stdout.write(self.style.SUCCESS(
            f'{prefix}Inventory update complete — '
            f'{stats["rows"]} rows, '
            f'{stats["matched"]} matched ({stats["matched_fuzzy"]} fuzzy), '
            f'{stats["updated"]} items updated, '
            f'{stats["stock_updated"]} stock rows updated, '
            f'{stats["skipped_no_match"]} no match, '
            f'{stats["skipped_ambiguous"]} ambiguous, '
            f'{stats["skipped_no_changes"]} unchanged.'
        ))
        if unmatched_samples:
            self.stdout.write('Unmatched samples:')
            for name in unmatched_samples:
                self.stdout.write(f'  - {name}')
